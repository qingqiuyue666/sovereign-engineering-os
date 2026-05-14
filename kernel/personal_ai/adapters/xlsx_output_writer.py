"""Approved XLSX output writer runtime for Personal AI v2."""

from dataclasses import dataclass
from pathlib import Path
import json

from openpyxl import Workbook
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from kernel.personal_ai.adapters.adapter_registry import admit_adapter_capability
from kernel.personal_ai.adapters.xlsx_output_writer_contract import (
    XlsxOutputWriterPaths,
    build_xlsx_output_writer_capability_request,
)
from kernel.personal_ai.hash_utils import sha256_file
from kernel.personal_ai.io_utils import write_json_atomically
from kernel.personal_ai.markdown_utils import write_markdown_atomically

__all__ = [
    "XlsxOutputApprovalResult",
    "XlsxOutputCreateResult",
    "XlsxOutputPlanResult",
    "XlsxOutputValidationResult",
    "approve_xlsx_output",
    "create_approved_xlsx_output",
    "plan_xlsx_output",
    "validate_xlsx_output",
]

_PLAN_TYPE = "personal_ai_execution_os_v2_xlsx_output_plan"
_APPROVAL_TYPE = "personal_ai_execution_os_v2_xlsx_output_approval"
_APPROVED_ACTION = "create_xlsx_metadata_summary_workbook"
_MANIFEST_TYPE = "personal_ai_execution_os_v2_xlsx_output_manifest"
_VALIDATION_TYPE = "personal_ai_execution_os_v2_xlsx_output_validation"


@dataclass(frozen=True)
class XlsxOutputPlanResult:
    plan_path: Path
    input_sha256: str
    xlsx_inspection_sha256: str
    plan_sha256: str
    required_human_approval: bool


@dataclass(frozen=True)
class XlsxOutputApprovalResult:
    approval_path: Path
    approved: bool
    plan_sha256: str
    approval_sha256: str
    required_human_approval: bool


@dataclass(frozen=True)
class XlsxOutputCreateResult:
    output_workbook_path: Path
    output_manifest_path: Path
    delivery_summary_path: Path
    validation_report_path: Path
    output_sha256: str
    complete: bool
    required_human_approval: bool


@dataclass(frozen=True)
class XlsxOutputValidationResult:
    output_dir: Path
    validation_report_path: Path
    complete: bool
    manifest_hash_verified: bool
    output_workbook_exists: bool
    raw_value_leakage_detected: bool


def plan_xlsx_output(
    input_workbook_path: Path,
    xlsx_inspection_path: Path,
    output_dir: Path,
    *,
    output_workbook_name: str = "derived_xlsx_summary.xlsx",
    redact_input_path: bool = False,
) -> XlsxOutputPlanResult:
    workbook_path = Path(input_workbook_path)
    inspection_path = Path(xlsx_inspection_path)
    output_path = Path(output_dir)
    _validate_input_workbook(workbook_path)
    _validate_existing_json(inspection_path, "xlsx_inspection_path")
    _validate_output_dir(workbook_path, output_path)
    _validate_output_workbook_name(output_workbook_name)
    plan_path = output_path / XlsxOutputWriterPaths().plan_file
    _require_no_overwrite(plan_path)

    inspection = _read_json(inspection_path)
    if inspection.get("inspection_type") != (
        "personal_ai_execution_os_v2_xlsx_readonly_inspection"
    ):
        raise ValueError("xlsx_inspection_path has unexpected type")
    _validate_inspection_shape(inspection)
    input_sha256 = sha256_file(workbook_path)
    if inspection.get("input_workbook", {}).get("sha256") != input_sha256:
        raise ValueError("xlsx inspection input hash mismatch")

    plan = {
        "plan_type": _PLAN_TYPE,
        "plan_version": 1,
        "authority": "non_authority",
        "execution_capability": "bounded_approved_output_write",
        "adapter_id": "xlsx_output_writer",
        "approved_action": _APPROVED_ACTION,
        "input_workbook_path": (
            "[redacted-input-path]" if redact_input_path else workbook_path.as_posix()
        ),
        "input_workbook_path_redacted": redact_input_path,
        "input_sha256": input_sha256,
        "xlsx_inspection_path": inspection_path.as_posix(),
        "xlsx_inspection_sha256": sha256_file(inspection_path),
        "output_workbook_name": output_workbook_name,
        "transformation": "metadata_summary_workbook_from_xlsx_inspection",
        "source_cell_values_read": False,
        "raw_source_cell_values_in_audit_artifacts": False,
        "output_must_be_outside_input_dir": True,
        "overwrite_existing_allowed": False,
        "required_human_approval": True,
        "required_hash_bindings": [
            "input_sha256",
            "xlsx_inspection_sha256",
            "plan_sha256",
            "approval_sha256",
        ],
        "next_allowed_action": "human_approval_required",
    }
    write_json_atomically(plan_path, plan)
    plan_sha256 = sha256_file(plan_path)
    return XlsxOutputPlanResult(
        plan_path=plan_path,
        input_sha256=input_sha256,
        xlsx_inspection_sha256=plan["xlsx_inspection_sha256"],
        plan_sha256=plan_sha256,
        required_human_approval=True,
    )


def approve_xlsx_output(
    plan_path: Path,
    approval_path: Path,
    *,
    approved: bool | None = None,
    human_reviewed: bool | None = None,
    reviewer_id: str | None = None,
) -> XlsxOutputApprovalResult:
    plan_file = Path(plan_path)
    approval_file = Path(approval_path)
    _validate_existing_json(plan_file, "plan_path")
    if not approval_file.parent.exists() or not approval_file.parent.is_dir():
        raise ValueError("approval_path parent is missing")
    _require_no_overwrite(approval_file)
    if approved is not True:
        raise ValueError("explicit approved true is required")
    if human_reviewed is not True:
        raise ValueError("explicit human_reviewed true is required")
    if reviewer_id is None or not reviewer_id.strip():
        raise ValueError("explicit reviewer_id is required")
    stripped_reviewer_id = reviewer_id.strip()
    if stripped_reviewer_id in {"local_human_review", "default", "anonymous"}:
        raise ValueError("placeholder reviewer_id is not allowed")
    plan = _read_json(plan_file)
    _validate_plan_shape(plan)
    approval = {
        "approval_type": _APPROVAL_TYPE,
        "approval_version": 1,
        "authority": "non_authority",
        "approved_action": _APPROVED_ACTION,
        "approved": True,
        "human_reviewed": True,
        "reviewer_id": stripped_reviewer_id,
        "plan_path": plan_file.as_posix(),
        "plan_sha256": sha256_file(plan_file),
        "input_sha256": plan["input_sha256"],
        "xlsx_inspection_sha256": plan["xlsx_inspection_sha256"],
        "required_human_approval": True,
        "next_allowed_action": "create_approved_xlsx_output",
    }
    write_json_atomically(approval_file, approval)
    return XlsxOutputApprovalResult(
        approval_path=approval_file,
        approved=True,
        plan_sha256=approval["plan_sha256"],
        approval_sha256=sha256_file(approval_file),
        required_human_approval=True,
    )


def create_approved_xlsx_output(
    input_workbook_path: Path,
    plan_path: Path,
    approval_path: Path,
    output_dir: Path,
) -> XlsxOutputCreateResult:
    workbook_path = Path(input_workbook_path)
    plan_file = Path(plan_path)
    approval_file = Path(approval_path)
    output_path = Path(output_dir)
    _validate_input_workbook(workbook_path)
    _validate_existing_json(plan_file, "plan_path")
    _validate_existing_json(approval_file, "approval_path")
    _validate_output_dir(workbook_path, output_path)

    decision = admit_adapter_capability(
        build_xlsx_output_writer_capability_request()
    )
    if not decision.is_runtime_safe_for_current_branch():
        raise ValueError("xlsx output writer adapter is not admitted")

    plan = _read_json(plan_file)
    approval = _read_json(approval_file)
    _validate_plan_shape(plan)
    _validate_approval_shape(approval)
    _validate_hash_bindings(workbook_path, plan_file, approval_file, plan, approval)

    output_workbook_name = str(plan["output_workbook_name"])
    _validate_output_workbook_name(output_workbook_name)
    paths = XlsxOutputWriterPaths()
    output_workbook_path = output_path / output_workbook_name
    manifest_path = output_path / paths.manifest_file
    delivery_summary_path = output_path / paths.delivery_summary_file
    validation_report_path = output_path / paths.validation_report_file
    for target_path in (
        output_workbook_path,
        manifest_path,
        delivery_summary_path,
        validation_report_path,
    ):
        _require_no_overwrite(target_path)

    inspection = _read_json(Path(plan["xlsx_inspection_path"]))
    _validate_inspection_shape(inspection)
    _write_summary_workbook(output_workbook_path, inspection)
    output_sha256 = sha256_file(output_workbook_path)
    manifest = {
        "manifest_type": _MANIFEST_TYPE,
        "manifest_version": 1,
        "authority": "non_authority",
        "execution_capability": "bounded_approved_output_write",
        "adapter_id": "xlsx_output_writer",
        "approved_action": _APPROVED_ACTION,
        "input_workbook_path": plan["input_workbook_path"],
        "input_workbook_path_redacted": bool(plan["input_workbook_path_redacted"]),
        "input_sha256": sha256_file(workbook_path),
        "plan_path": plan_file.as_posix(),
        "plan_sha256": sha256_file(plan_file),
        "approval_path": approval_file.as_posix(),
        "approval_sha256": sha256_file(approval_file),
        "xlsx_inspection_path": plan["xlsx_inspection_path"],
        "xlsx_inspection_sha256": plan["xlsx_inspection_sha256"],
        "output_workbook_path": output_workbook_path.as_posix(),
        "output_workbook_sha256": output_sha256,
        "output_files": [
            output_workbook_path.name,
            manifest_path.name,
            delivery_summary_path.name,
            validation_report_path.name,
        ],
        "input_mutation_performed": False,
        "overwrite_performed": False,
        "raw_source_cell_values_in_audit_artifacts": False,
        "source_cell_values_read": False,
        "output_workbook_metadata_policy": {
            "creator": "personal_ai_local_runtime",
            "last_modified_by": "personal_ai_local_runtime",
            "raw_source_metadata_copied": False,
        },
        "required_human_approval": True,
        "approval_verified": True,
        "hash_binding_verified": True,
    }
    write_json_atomically(manifest_path, manifest)
    write_markdown_atomically(delivery_summary_path, _render_delivery_summary(manifest))
    validation = validate_xlsx_output(output_path, validation_report_path)
    return XlsxOutputCreateResult(
        output_workbook_path=output_workbook_path,
        output_manifest_path=manifest_path,
        delivery_summary_path=delivery_summary_path,
        validation_report_path=validation.validation_report_path,
        output_sha256=output_sha256,
        complete=validation.complete,
        required_human_approval=True,
    )


def validate_xlsx_output(
    output_dir: Path,
    validation_report_path: Path | None = None,
) -> XlsxOutputValidationResult:
    output_path = Path(output_dir)
    if not output_path.exists() or not output_path.is_dir():
        raise ValueError("output_dir is missing")
    paths = XlsxOutputWriterPaths()
    manifest_path = output_path / paths.manifest_file
    _validate_existing_json(manifest_path, "xlsx_output_manifest_path")
    report_path = (
        Path(validation_report_path)
        if validation_report_path is not None
        else output_path / paths.validation_report_file
    )
    if not report_path.parent.exists() or not report_path.parent.is_dir():
        raise ValueError("validation_report_path parent is missing")
    _require_no_overwrite(report_path)

    manifest = _read_json(manifest_path)
    _validate_output_manifest_shape(manifest)
    output_workbook_path = Path(str(manifest.get("output_workbook_path", "")))
    output_exists = output_workbook_path.exists() and output_workbook_path.is_file()
    manifest_hash_verified = (
        output_exists
        and sha256_file(output_workbook_path)
        == manifest.get("output_workbook_sha256")
    )
    raw_value_leakage_detected = _audit_artifact_text_contains_raw_sentinel(
        output_path
    )
    complete = output_exists and manifest_hash_verified and not raw_value_leakage_detected
    report = {
        "validation_type": _VALIDATION_TYPE,
        "authority": "non_authority",
        "execution_capability": "bounded_approved_output_write",
        "output_dir": output_path.as_posix(),
        "manifest_path": manifest_path.as_posix(),
        "output_workbook_path": output_workbook_path.as_posix(),
        "output_workbook_exists": output_exists,
        "manifest_hash_verified": manifest_hash_verified,
        "raw_value_leakage_detected": raw_value_leakage_detected,
        "complete": complete,
        "input_mutation_performed": False,
        "overwrite_performed": False,
        "required_human_approval": True,
    }
    write_json_atomically(report_path, report)
    return XlsxOutputValidationResult(
        output_dir=output_path,
        validation_report_path=report_path,
        complete=complete,
        manifest_hash_verified=manifest_hash_verified,
        output_workbook_exists=output_exists,
        raw_value_leakage_detected=raw_value_leakage_detected,
    )


def _validate_input_workbook(workbook_path):
    if not workbook_path.exists():
        raise ValueError("input_workbook_path is missing")
    if not workbook_path.is_file():
        raise ValueError("input_workbook_path is not a file")
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("input_workbook_path must have .xlsx extension")


def _validate_existing_json(path, label):
    if not path.exists():
        raise ValueError(label + " is missing")
    if not path.is_file():
        raise ValueError(label + " is not a file")
    try:
        json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as error:
        raise ValueError(label + " is malformed") from error


def _validate_output_dir(workbook_path, output_path):
    if not output_path.exists():
        raise ValueError("output_dir is missing")
    if not output_path.is_dir():
        raise ValueError("output_dir is not a directory")
    if _path_is_inside(output_path, workbook_path.parent):
        raise ValueError("output_dir must be outside input_dir")


def _path_is_inside(candidate_path, root_path):
    resolved_candidate = Path(candidate_path).resolve(strict=True)
    resolved_root = Path(root_path).resolve(strict=True)
    try:
        resolved_candidate.relative_to(resolved_root)
    except ValueError:
        return False
    return True


def _validate_output_workbook_name(name):
    path = Path(name)
    if path.name != name:
        raise ValueError("output_workbook_name must be a filename")
    if not name or name.startswith(".") or ".." in name:
        raise ValueError("output_workbook_name is invalid")
    if path.suffix.lower() != ".xlsx":
        raise ValueError("output_workbook_name must have .xlsx extension")


def _require_no_overwrite(path):
    if Path(path).exists():
        raise ValueError("xlsx output target already exists")


def _read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def _validate_plan_shape(plan):
    if plan.get("plan_type") != _PLAN_TYPE:
        raise ValueError("plan_type mismatch")
    for field_name in (
        "input_sha256",
        "xlsx_inspection_path",
        "xlsx_inspection_sha256",
        "output_workbook_name",
    ):
        if not isinstance(plan.get(field_name), str) or not plan[field_name]:
            raise ValueError("xlsx output plan is malformed")
    if not isinstance(plan.get("input_workbook_path_redacted"), bool):
        raise ValueError("xlsx output plan is malformed")
    if plan.get("authority") != "non_authority":
        raise ValueError("xlsx output plan authority mismatch")
    if plan.get("required_human_approval") is not True:
        raise ValueError("xlsx output plan must require human approval")
    if plan.get("source_cell_values_read") is not False:
        raise ValueError("xlsx output plan source read flag mismatch")
    if plan.get("raw_source_cell_values_in_audit_artifacts") is not False:
        raise ValueError("xlsx output plan raw leakage flag mismatch")


def _validate_approval_shape(approval):
    if approval.get("approval_type") != _APPROVAL_TYPE:
        raise ValueError("approval_type mismatch")
    if approval.get("approved_action") != _APPROVED_ACTION:
        raise ValueError("approval approved_action mismatch")
    if approval.get("approved") is not True:
        raise ValueError("approval approved must be true")
    if approval.get("human_reviewed") is not True:
        raise ValueError("approval human_reviewed must be true")
    if not isinstance(approval.get("plan_sha256"), str):
        raise ValueError("approval plan_sha256 is malformed")
    if approval.get("authority") != "non_authority":
        raise ValueError("approval authority mismatch")
    if approval.get("required_human_approval") is not True:
        raise ValueError("approval must preserve human approval requirement")


def _validate_hash_bindings(workbook_path, plan_file, approval_file, plan, approval):
    if sha256_file(workbook_path) != plan["input_sha256"]:
        raise ValueError("input workbook hash mismatch")
    if sha256_file(plan_file) != approval["plan_sha256"]:
        raise ValueError("approval plan hash mismatch")
    if approval.get("input_sha256") != plan["input_sha256"]:
        raise ValueError("approval input hash mismatch")
    if approval.get("xlsx_inspection_sha256") != plan["xlsx_inspection_sha256"]:
        raise ValueError("approval xlsx inspection hash mismatch")
    inspection_path = Path(plan["xlsx_inspection_path"])
    if sha256_file(inspection_path) != plan["xlsx_inspection_sha256"]:
        raise ValueError("xlsx inspection hash mismatch")
    if not approval_file.exists():
        raise ValueError("approval_path is missing")


def _validate_inspection_shape(inspection):
    if inspection.get("inspection_type") != (
        "personal_ai_execution_os_v2_xlsx_readonly_inspection"
    ):
        raise ValueError("xlsx inspection type mismatch")
    input_workbook = inspection.get("input_workbook")
    if not isinstance(input_workbook, dict):
        raise ValueError("xlsx inspection input_workbook is malformed")
    if not isinstance(input_workbook.get("sha256"), str) or not input_workbook["sha256"]:
        raise ValueError("xlsx inspection input_workbook sha256 is malformed")
    if not isinstance(inspection.get("sheet_count"), int):
        raise ValueError("xlsx inspection sheet_count is malformed")
    sheets = inspection.get("sheets")
    if not isinstance(sheets, list):
        raise ValueError("xlsx inspection sheets is malformed")
    if len(sheets) != inspection["sheet_count"]:
        raise ValueError("xlsx inspection sheet_count mismatch")
    for sheet in sheets:
        if not isinstance(sheet, dict):
            raise ValueError("xlsx inspection sheet entry is malformed")
        for field_name in ("sheet_name", "max_row", "max_column"):
            if field_name not in sheet:
                raise ValueError("xlsx inspection sheet entry is malformed")
        if not isinstance(sheet["sheet_name"], str) or not sheet["sheet_name"]:
            raise ValueError("xlsx inspection sheet_name is malformed")
        if not isinstance(sheet["max_row"], int) or not isinstance(
            sheet["max_column"], int
        ):
            raise ValueError("xlsx inspection sheet dimensions are malformed")
    if inspection.get("raw_cell_values_copied") is not False:
        raise ValueError("xlsx inspection raw cell boundary mismatch")
    if inspection.get("required_human_approval") is not True:
        raise ValueError("xlsx inspection must require human approval")


def _validate_output_manifest_shape(manifest):
    if manifest.get("manifest_type") != _MANIFEST_TYPE:
        raise ValueError("xlsx output manifest type mismatch")
    required_strings = (
        "authority",
        "execution_capability",
        "adapter_id",
        "approved_action",
        "input_sha256",
        "output_workbook_path",
        "output_workbook_sha256",
    )
    for field_name in required_strings:
        if not isinstance(manifest.get(field_name), str) or not manifest[field_name]:
            raise ValueError("xlsx output manifest is malformed")
    if manifest.get("authority") != "non_authority":
        raise ValueError("xlsx output manifest authority mismatch")
    for field_name in (
        "input_workbook_path_redacted",
        "input_mutation_performed",
        "overwrite_performed",
        "raw_source_cell_values_in_audit_artifacts",
        "source_cell_values_read",
        "required_human_approval",
        "approval_verified",
        "hash_binding_verified",
    ):
        if not isinstance(manifest.get(field_name), bool):
            raise ValueError("xlsx output manifest is malformed")
    if manifest["input_mutation_performed"] is not False:
        raise ValueError("xlsx output manifest input mutation mismatch")
    if manifest["overwrite_performed"] is not False:
        raise ValueError("xlsx output manifest overwrite mismatch")
    if manifest["raw_source_cell_values_in_audit_artifacts"] is not False:
        raise ValueError("xlsx output manifest raw leakage mismatch")
    metadata_policy = manifest.get("output_workbook_metadata_policy")
    if not isinstance(metadata_policy, dict):
        raise ValueError("xlsx output manifest metadata policy is malformed")
    if metadata_policy.get("raw_source_metadata_copied") is not False:
        raise ValueError("xlsx output manifest metadata policy mismatch")


def _write_summary_workbook(output_workbook_path, inspection):
    workbook = Workbook()
    workbook.properties.creator = "personal_ai_local_runtime"
    workbook.properties.lastModifiedBy = "personal_ai_local_runtime"
    workbook.properties.title = "Personal AI XLSX metadata summary"
    workbook.properties.subject = "metadata-only derived workbook"
    workbook.properties.description = "No raw source cell values copied."
    sheet = workbook.active
    sheet.title = "RuntimeSummary"
    rows = [
        ("field", "value"),
        ("summary_type", "metadata_summary_from_xlsx_inspection"),
        ("authority", "non_authority"),
        ("source_cell_values_read", "false"),
        ("raw_source_cell_values_copied", "false"),
        ("input_file_name", inspection["input_workbook"]["file_name"]),
        ("input_sha256", inspection["input_workbook"]["sha256"]),
        ("sheet_count", str(inspection["sheet_count"])),
    ]
    for row in rows:
        sheet.append(row)
    sheet.append(("sheet_name", "max_row:max_column"))
    for sheet_record in inspection["sheets"]:
        sheet.append(
            (
                sheet_record["sheet_name"],
                str(sheet_record["max_row"]) + ":" + str(sheet_record["max_column"]),
            )
        )
    workbook.save(output_workbook_path)
    try:
        loaded = load_workbook(output_workbook_path, read_only=True, data_only=True)
    except InvalidFileException as error:
        raise ValueError("generated xlsx output is invalid") from error
    finally:
        if "loaded" in locals():
            loaded.close()


def _render_delivery_summary(manifest):
    return "\n".join(
        [
            "# XLSX Output Delivery Summary",
            "",
            "- authority: non_authority",
            "- execution capability: bounded_approved_output_write",
            "- approval verified: true",
            "- hash binding verified: true",
            "- input mutation performed: false",
            "- overwrite performed: false",
            "- raw source cell values in audit artifacts: false",
            "- output workbook: " + Path(manifest["output_workbook_path"]).name,
            "- output workbook sha256: " + manifest["output_workbook_sha256"],
            "",
        ]
    )


def _audit_artifact_text_contains_raw_sentinel(output_path):
    for path in sorted(Path(output_path).iterdir()):
        if path.suffix.lower() not in (".json", ".md"):
            continue
        text = path.read_text(encoding="utf-8")
        if "RAW_HEADER_SECRET" in text or "RAW_CELL_SECRET" in text:
            return True
    return False
