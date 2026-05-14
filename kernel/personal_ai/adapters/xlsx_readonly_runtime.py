"""Bounded readonly XLSX inspection runtime for Personal AI v2."""

from dataclasses import dataclass
from pathlib import Path
from zipfile import BadZipFile
import hashlib

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from kernel.personal_ai.adapters.adapter_registry import admit_adapter_capability
from kernel.personal_ai.adapters.xlsx_adapter_contract import (
    XlsxReadonlyLimits,
    build_xlsx_readonly_capability_request,
)
from kernel.personal_ai.hash_utils import sha256_file
from kernel.personal_ai.io_utils import write_json_atomically
from kernel.personal_ai.markdown_utils import write_markdown_atomically

__all__ = [
    "XlsxReadonlyInspectionResult",
    "inspect_xlsx_readonly",
]

_INSPECTION_FILE = "xlsx_inspection.json"
_SUMMARY_FILE = "xlsx_inspection_summary.md"


@dataclass(frozen=True)
class XlsxReadonlyInspectionResult:
    input_workbook_path: Path
    output_dir: Path
    xlsx_inspection_path: Path
    xlsx_inspection_summary_path: Path
    input_sha256: str
    sheet_count: int
    required_human_approval: bool


def inspect_xlsx_readonly(
    input_workbook_path: Path,
    output_dir: Path,
    *,
    limits: XlsxReadonlyLimits | None = None,
    redact_sheet_names: bool = False,
    redact_input_path: bool = False,
) -> XlsxReadonlyInspectionResult:
    workbook_path = Path(input_workbook_path)
    output_path = Path(output_dir)
    effective_limits = limits or XlsxReadonlyLimits()
    _validate_paths(workbook_path, output_path, effective_limits)

    decision = admit_adapter_capability(build_xlsx_readonly_capability_request())
    if not decision.is_runtime_safe_for_current_branch():
        raise ValueError("xlsx readonly adapter is not admitted")

    inspection_path = output_path / _INSPECTION_FILE
    summary_path = output_path / _SUMMARY_FILE
    _require_no_overwrite(inspection_path)
    _require_no_overwrite(summary_path)

    input_sha256 = sha256_file(workbook_path)
    try:
        workbook = load_workbook(
            workbook_path,
            read_only=True,
            data_only=False,
            keep_links=False,
        )
    except (InvalidFileException, BadZipFile, EOFError, KeyError, OSError) as error:
        raise ValueError("input_workbook_path is not a valid xlsx file") from error

    warnings = []
    try:
        sheets = [
            _inspect_sheet(
                sheet,
                effective_limits,
                warnings,
                sheet_index=sheet_index,
                redact_sheet_names=redact_sheet_names,
            )
            for sheet_index, sheet in enumerate(workbook.worksheets, start=1)
        ]
    finally:
        workbook.close()

    input_workbook_record = {
        "file_name": (
            "[redacted-input-file-name]"
            if redact_input_path
            else workbook_path.name
        ),
        "path": "[redacted-input-path]" if redact_input_path else workbook_path.as_posix(),
        "sha256": input_sha256,
        "path_redacted": redact_input_path,
        "file_name_redacted": redact_input_path,
    }
    payload = {
        "inspection_type": "personal_ai_execution_os_v2_xlsx_readonly_inspection",
        "authority": "non_authority",
        "execution_capability": "bounded_local_readonly_runtime",
        "input_workbook": input_workbook_record,
        "input_mutation_performed": False,
        "output_overwrite_performed": False,
        "raw_cell_values_copied": False,
        "raw_sheet_names_copied": not redact_sheet_names,
        "required_human_approval": True,
        "adapter_id": "xlsx_readonly_runtime",
        "sheet_count": len(sheets),
        "sheet_names": [sheet["sheet_name"] for sheet in sheets],
        "redaction": {
            "input_path_redacted": redact_input_path,
            "input_file_name_redacted": redact_input_path,
            "sheet_names_redacted": redact_sheet_names,
        },
        "sheets": sheets,
        "extraction_limits": {
            "max_header_rows": effective_limits.max_header_rows,
            "max_header_columns": effective_limits.max_header_columns,
            "max_formula_scan_cells": effective_limits.max_formula_scan_cells,
            "max_style_scan_cells": effective_limits.max_style_scan_cells,
            "max_workbook_bytes": effective_limits.max_workbook_bytes,
            "header_preview_raw_values": False,
        },
        "warnings": sorted(set(warnings)),
        "forbidden_actions": [
            "modify_input_files",
            "overwrite_existing_outputs",
            "copy_full_raw_workbook_values",
            "call_network",
            "run_subprocess",
            "control_external_tools",
        ],
        "boundaries": {
            "network_allowed": False,
            "subprocess_allowed": False,
            "external_tool_control_allowed": False,
            "input_mutation_allowed": False,
            "overwrite_existing_allowed": False,
            "raw_value_copy_allowed": False,
            "requires_manifest": True,
            "requires_evidence_capture": True,
            "requires_output_hashing": True,
        },
        "next_allowed_action": "human_review_only",
    }
    write_json_atomically(inspection_path, payload)
    write_markdown_atomically(summary_path, _render_summary(payload))

    return XlsxReadonlyInspectionResult(
        input_workbook_path=workbook_path,
        output_dir=output_path,
        xlsx_inspection_path=inspection_path,
        xlsx_inspection_summary_path=summary_path,
        input_sha256=input_sha256,
        sheet_count=len(sheets),
        required_human_approval=True,
    )


def _validate_paths(workbook_path, output_path, limits):
    limits.validate()
    if not workbook_path.exists():
        raise ValueError("input_workbook_path is missing")
    if not workbook_path.is_file():
        raise ValueError("input_workbook_path is not a file")
    if workbook_path.is_symlink():
        raise ValueError("input_workbook_path must not be a symlink")
    if workbook_path.suffix.lower() != ".xlsx":
        raise ValueError("input_workbook_path must have .xlsx extension")
    if workbook_path.stat().st_size > limits.max_workbook_bytes:
        raise ValueError("input_workbook_path exceeds max_workbook_bytes")
    if not output_path.exists():
        raise ValueError("output_dir is missing")
    if not output_path.is_dir():
        raise ValueError("output_dir is not a directory")
    if output_path.is_symlink():
        raise ValueError("output_dir must not be a symlink")
    if _path_is_inside(output_path, workbook_path.parent):
        raise ValueError("output_dir must be outside input_dir")


def _require_no_overwrite(path):
    if Path(path).exists():
        raise ValueError("xlsx inspection output already exists")


def _path_is_inside(candidate_path, root_path):
    resolved_candidate = Path(candidate_path).resolve(strict=True)
    resolved_root = Path(root_path).resolve(strict=True)
    try:
        resolved_candidate.relative_to(resolved_root)
    except ValueError:
        return False
    return True


def _inspect_sheet(sheet, limits, warnings, *, sheet_index, redact_sheet_names):
    formula_stats = _scan_formula_stats(sheet, limits.max_formula_scan_cells)
    style_stats = _scan_style_stats(sheet, limits.max_style_scan_cells)
    merged_count = _merged_range_count(sheet, warnings)
    raw_sheet_name = sheet.title
    safe_sheet_name = (
        _redacted_sheet_label(sheet_index) if redact_sheet_names else raw_sheet_name
    )
    return {
        "sheet_name": safe_sheet_name,
        "sheet_index": sheet_index,
        "sheet_name_redacted": redact_sheet_names,
        "raw_sheet_name_included": not redact_sheet_names,
        "sheet_name_sha256": hashlib.sha256(
            raw_sheet_name.encode("utf-8")
        ).hexdigest(),
        "max_row": int(sheet.max_row or 0),
        "max_column": int(sheet.max_column or 0),
        "merged_cell_range_count": merged_count,
        "formula_presence": formula_stats["formula_presence"],
        "formula_count_observed": formula_stats["formula_count_observed"],
        "formula_scan_truncated": formula_stats["formula_scan_truncated"],
        "style_presence": style_stats["style_presence"],
        "style_scan_truncated": style_stats["style_scan_truncated"],
        "header_preview": _header_preview(sheet, limits),
    }


def _redacted_sheet_label(sheet_index):
    return "sheet_" + str(sheet_index)


def _merged_range_count(sheet, warnings):
    merged_cells = getattr(sheet, "merged_cells", None)
    if merged_cells is None:
        warnings.append("merged_cell_ranges_unavailable_in_readonly_mode")
        return None
    ranges = getattr(merged_cells, "ranges", None)
    if ranges is None:
        warnings.append("merged_cell_ranges_unavailable_in_readonly_mode")
        return None
    return len(tuple(ranges))


def _scan_formula_stats(sheet, max_cells):
    scanned = 0
    formula_count = 0
    truncated = False
    for row in sheet.iter_rows():
        for cell in row:
            scanned += 1
            if _cell_has_formula(cell):
                formula_count += 1
            if scanned >= max_cells:
                truncated = True
                return {
                    "formula_presence": formula_count > 0,
                    "formula_count_observed": formula_count,
                    "formula_scan_truncated": truncated,
                }
    return {
        "formula_presence": formula_count > 0,
        "formula_count_observed": formula_count,
        "formula_scan_truncated": truncated,
    }


def _scan_style_stats(sheet, max_cells):
    scanned = 0
    style_presence = False
    truncated = False
    for row in sheet.iter_rows():
        for cell in row:
            scanned += 1
            style_id = getattr(cell, "style_id", 0)
            if style_id not in (None, 0):
                style_presence = True
            if scanned >= max_cells:
                truncated = True
                return {
                    "style_presence": style_presence,
                    "style_scan_truncated": truncated,
                }
    return {
        "style_presence": style_presence,
        "style_scan_truncated": truncated,
    }


def _header_preview(sheet, limits):
    preview = []
    for row_index, row in enumerate(
        sheet.iter_rows(
            max_row=limits.max_header_rows,
            max_col=limits.max_header_columns,
        ),
        start=1,
    ):
        preview_row = []
        for column_index, cell in enumerate(row, start=1):
            value = cell.value
            preview_row.append(
                {
                    "row": row_index,
                    "column": column_index,
                    "value_type": _safe_value_type(value),
                    "is_blank": value in (None, ""),
                    "text_length": len(value) if isinstance(value, str) else None,
                    "has_formula": _cell_has_formula(cell),
                    "raw_value_included": False,
                }
            )
        preview.append(preview_row)
    return preview


def _cell_has_formula(cell):
    if getattr(cell, "data_type", None) == "f":
        return True
    value = getattr(cell, "value", None)
    return isinstance(value, str) and value.startswith("=")


def _safe_value_type(value):
    if value is None:
        return "blank"
    if isinstance(value, bool):
        return "bool"
    if isinstance(value, int):
        return "int"
    if isinstance(value, float):
        return "float"
    if isinstance(value, str):
        return "formula" if value.startswith("=") else "text"
    return type(value).__name__


def _render_summary(payload):
    lines = [
        "# XLSX Readonly Inspection Summary",
        "",
        "## Authority",
        "",
        "- authority: non_authority",
        "- execution capability: bounded_local_readonly_runtime",
        "- required human approval: true",
        "- input mutation performed: false",
        "- raw cell values copied: false",
        "",
        "## Workbook",
        "",
        "- file name: " + payload["input_workbook"]["file_name"],
        "- input sha256: " + payload["input_workbook"]["sha256"],
        "- sheet count: " + str(payload["sheet_count"]),
        "",
        "## Sheets",
        "",
    ]
    for sheet in payload["sheets"]:
        lines.extend(
            [
                "- "
                + sheet["sheet_name"]
                + ": rows="
                + str(sheet["max_row"])
                + ", columns="
                + str(sheet["max_column"])
                + ", formulas="
                + str(sheet["formula_count_observed"]),
            ]
        )
    lines.extend(
        [
            "",
            "## Limits",
            "",
            "- header preview raw values: false",
            "- no full workbook cell contents copied",
            "",
        ]
    )
    return "\n".join(lines)
