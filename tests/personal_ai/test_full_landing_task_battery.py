import json
import tempfile
import time
import unittest
from pathlib import Path

from openpyxl import Workbook

from kernel.personal_ai.adapters.blender_controlled_runtime import (
    run_blender_controlled_fixture,
)
from kernel.personal_ai.adapters.browser_fixture_runtime import (
    run_browser_fixture_from_actions_file,
)
from kernel.personal_ai.adapters.comfyui_controlled_runtime import (
    run_comfyui_controlled_fixture,
)
from kernel.personal_ai.adapters.model_adapter_contract import ModelFixtureSchema
from kernel.personal_ai.adapters.model_typed_schema_runtime import (
    run_model_fixture,
    write_model_provider_request,
)
from kernel.personal_ai.adapters.xlsx_readonly_runtime import inspect_xlsx_readonly
from kernel.personal_ai.hash_utils import sha256_file
from kernel.personal_ai.io_utils import write_json_atomically
from kernel.personal_ai.local_launcher import (
    run_browser_fixture_launcher,
    run_local_office_launcher,
    run_model_fixture_launcher,
    run_runtime_delivery_validation_launcher,
)
from kernel.personal_ai.runtime_delivery_package import (
    build_runtime_delivery_package,
    validate_runtime_delivery_package,
)
from kernel.personal_ai.task_graph import run_local_task_graph_fixture


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


class FullLandingTaskBatteryTests(unittest.TestCase):
    def make_root(self):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        return Path(temp_dir.name)

    def write_workbook(self, path, *, rows=12):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Finance Ops"
        sheet["A1"] = "RAW_BATTERY_HEADER_SENTINEL"
        sheet["B1"] = "Amount"
        for row_index in range(2, rows + 2):
            sheet.cell(row=row_index, column=1).value = (
                "RAW_BATTERY_CELL_SENTINEL_" + str(row_index)
            )
            sheet.cell(row=row_index, column=2).value = row_index * 10
        sheet["C2"] = "=SUM(B2:B4)"
        audit = workbook.create_sheet("Audit Trail")
        audit["A1"] = "Reviewer"
        audit["B1"] = "Status"
        workbook.save(path)

    def write_browser_fixture(self, fixture_path, actions_path):
        fixture_path.write_text(
            "<html><head><title>Battery Fixture</title></head><body>"
            "<a href='https://allowed.example/reference'>Allowed</a>"
            "<input name='query' data-seos-allowed-field='true'>"
            "<button id='inspect' data-seos-allowed-button='true'>Inspect</button>"
            "</body></html>",
            encoding="utf-8",
        )
        write_json_atomically(
            actions_path,
            {
                "actions": [
                    {"action": "open_local_fixture"},
                    {"action": "inspect_title"},
                    {
                        "action": "fill_allowed_field",
                        "field_name": "query",
                        "value": "local fixture value",
                    },
                    {"action": "click_allowed_button", "button_id": "inspect"},
                ],
                "policy": {
                    "allowed_domains": ["allowed.example"],
                    "timeout_seconds": 5,
                },
            },
        )

    def test_realistic_xlsx_malformed_and_large_workbook_battery(self):
        root = self.make_root()
        input_dir = root / "input"
        output_dir = root / "xlsx-output"
        large_output_dir = root / "large-output"
        damaged_output_dir = root / "damaged-output"
        for path in (input_dir, output_dir, large_output_dir, damaged_output_dir):
            path.mkdir()
        workbook_path = input_dir / "realistic.xlsx"
        self.write_workbook(workbook_path)
        before_hash = sha256_file(workbook_path)

        result = inspect_xlsx_readonly(
            workbook_path,
            output_dir,
            redact_sheet_names=True,
            redact_input_path=True,
        )
        inspection_text = result.xlsx_inspection_path.read_text(encoding="utf-8")

        self.assertEqual(before_hash, sha256_file(workbook_path))
        self.assertNotIn("RAW_BATTERY_HEADER_SENTINEL", inspection_text)
        self.assertNotIn("RAW_BATTERY_CELL_SENTINEL", inspection_text)
        self.assertNotIn("Finance Ops", inspection_text)
        self.assertTrue(read_json(result.xlsx_inspection_path)["redaction"]["sheet_names_redacted"])

        large_workbook_path = input_dir / "large-smoke.xlsx"
        self.write_workbook(large_workbook_path, rows=1200)
        start = time.perf_counter()
        large_result = inspect_xlsx_readonly(large_workbook_path, large_output_dir)
        duration_seconds = time.perf_counter() - start

        self.assertEqual(large_result.sheet_count, 2)
        self.assertLess(duration_seconds, 10.0)

        damaged_path = input_dir / "damaged.xlsx"
        damaged_path.write_bytes(b"not a valid workbook")
        with self.assertRaisesRegex(ValueError, "not a valid xlsx"):
            inspect_xlsx_readonly(damaged_path, damaged_output_dir)

    def test_boundary_violation_fixtures_and_failure_quarantine_battery(self):
        root = self.make_root()
        comfy_dir = root / "comfy"
        blender_dir = root / "blender"
        model_dir = root / "model"
        for path in (comfy_dir, blender_dir, model_dir):
            path.mkdir()

        workflow_path = root / "workflow.json"
        write_json_atomically(
            workflow_path,
            {"nodes": [{"id": "bad", "type": "PythonScript", "inputs": {}}]},
        )
        comfy_result = run_comfyui_controlled_fixture(workflow_path, comfy_dir)
        comfy_failure = read_json(comfy_result.failure_bundle_path)
        self.assertFalse(comfy_result.success)
        self.assertFalse(comfy_failure["arbitrary_node_execution_performed"])

        scene_path = root / "scene.blend"
        scene_path.write_bytes(b"BLENDER_BOUNDARY_SCENE")
        plan_path = root / "blender_plan.json"
        write_json_atomically(
            plan_path,
            {
                "operations": [
                    {
                        "operation": "export_glb_preview",
                        "parameters": {"target_path": scene_path.as_posix()},
                    }
                ]
            },
        )
        blender_result = run_blender_controlled_fixture(
            scene_path,
            plan_path,
            blender_dir,
        )
        blender_failure = read_json(blender_result.failure_bundle_path)
        self.assertFalse(blender_result.success)
        self.assertFalse(blender_failure["source_asset_overwrite_performed"])
        self.assertFalse(blender_failure["real_blender_runtime_called"])

        input_artifact = root / "model_input.json"
        request_path = root / "model_request.json"
        write_json_atomically(input_artifact, {"route_type": "spreadsheet_review"})
        write_model_provider_request(
            input_artifact,
            request_path,
            schema_name=ModelFixtureSchema.JOB_ROUTE_CLASSIFICATION,
        )
        request = read_json(request_path)
        request["fixture_response"] = {
            "schema_name": ModelFixtureSchema.JOB_ROUTE_CLASSIFICATION,
            "authority": "non_authority",
            "can_grant_authority": True,
            "can_edit_files": False,
            "can_call_tools": False,
        }
        write_json_atomically(request_path, request)

        model_result = run_model_fixture(request_path, model_dir)
        model_failure = read_json(model_result.failure_bundle_path)
        self.assertFalse(model_result.success)
        self.assertTrue(model_failure["invalid_output_quarantined"])
        self.assertFalse(model_failure["network_used"])

    def test_browser_allow_deny_and_delivery_tamper_battery(self):
        root = self.make_root()
        input_dir = root / "input"
        browser_dir = root / "browser"
        denied_dir = root / "browser-denied"
        runtime_dir = root / "runtime"
        package_root = root / "packages"
        validation_dir = root / "validation"
        for path in (
            input_dir,
            browser_dir,
            denied_dir,
            runtime_dir,
            package_root,
            validation_dir,
        ):
            path.mkdir()
        fixture_path = input_dir / "fixture.html"
        actions_path = input_dir / "actions.json"
        self.write_browser_fixture(fixture_path, actions_path)

        browser_result = run_browser_fixture_from_actions_file(
            fixture_path,
            actions_path,
            browser_dir,
        )
        self.assertEqual(browser_result.action_count, 4)
        self.assertFalse(
            read_json(browser_result.evidence_manifest_path)[
                "real_screenshot_captured"
            ]
        )

        denied_fixture = input_dir / "denied.html"
        denied_fixture.write_text(
            "<html><body><a href='https://denied.example'>Denied</a></body></html>",
            encoding="utf-8",
        )
        denied_actions = input_dir / "denied_actions.json"
        write_json_atomically(
            denied_actions,
            {"actions": [{"action": "inspect_links"}], "policy": {}},
        )
        with self.assertRaisesRegex(ValueError, "not in domain allowlist"):
            run_browser_fixture_from_actions_file(
                denied_fixture,
                denied_actions,
                denied_dir,
            )

        write_json_atomically(
            runtime_dir / "browser_action_log.json",
            {"log_type": "personal_ai_execution_os_v2_browser_action_log"},
        )
        package = build_runtime_delivery_package(
            runtime_dir,
            package_root,
            package_id="tamper-delivery",
        )
        (package.package_dir / "browser_action_log.json").write_text(
            "{\"tampered\": true}\n",
            encoding="utf-8",
        )
        validation = validate_runtime_delivery_package(
            package.package_dir,
            validation_dir / "runtime_delivery_validation.json",
        )
        validation_payload = read_json(validation.runtime_delivery_validation_path)
        self.assertFalse(validation.complete)
        self.assertIn("browser_action_log", validation_payload["hash_mismatches"])

    def test_full_local_e2e_battery_via_launchers_and_task_graph(self):
        root = self.make_root()
        input_dir = root / "input"
        office_dir = root / "office"
        model_dir = root / "model"
        browser_dir = root / "browser"
        runtime_dir = root / "runtime"
        package_root = root / "packages"
        delivery_validation_dir = root / "delivery-validation"
        graph_dir = root / "graph"
        for path in (
            input_dir,
            office_dir,
            model_dir,
            browser_dir,
            runtime_dir,
            package_root,
            delivery_validation_dir,
            graph_dir,
        ):
            path.mkdir()
        workbook_path = input_dir / "source.xlsx"
        self.write_workbook(workbook_path)
        model_input_path = input_dir / "model_input.json"
        write_json_atomically(
            model_input_path,
            {"route_type": "spreadsheet_review", "recommended_processor_lane": "human_review_only"},
        )
        fixture_path = input_dir / "fixture.html"
        actions_path = input_dir / "actions.json"
        self.write_browser_fixture(fixture_path, actions_path)
        workbook_hash_before = sha256_file(workbook_path)
        fixture_hash_before = sha256_file(fixture_path)

        office = run_local_office_launcher(workbook_path, office_dir)
        model = run_model_fixture_launcher(
            model_input_path,
            model_dir,
            schema_name=ModelFixtureSchema.JOB_ROUTE_CLASSIFICATION,
        )
        browser = run_browser_fixture_launcher(fixture_path, actions_path, browser_dir)
        write_json_atomically(
            runtime_dir / "model_inference_artifact.json",
            read_json(model.payload["model_inference_artifact_path"]),
        )
        write_json_atomically(
            runtime_dir / "browser_action_log.json",
            read_json(browser.payload["browser_action_log_path"]),
        )
        package = build_runtime_delivery_package(
            runtime_dir,
            package_root,
            package_id="full-e2e-delivery",
            input_dir=input_dir,
        )
        delivery = run_runtime_delivery_validation_launcher(
            package.package_dir,
            delivery_validation_dir,
        )
        graph_path = root / "task_graph.json"
        write_json_atomically(
            graph_path,
            {
                "graph_type": "personal_ai_execution_os_unified_task_graph_v1",
                "graph_id": "full-e2e",
                "authority": "non_authority",
                "required_human_approval": True,
                "nodes": [
                    {
                        "node_id": "model_contract",
                        "adapter_id": "mock_model_typed_schema_runtime",
                        "capability": "classify_local_job_package",
                        "depends_on": [],
                        "approval_checkpoint_required": True,
                        "inputs": {},
                    },
                    {
                        "node_id": "delivery_validation",
                        "adapter_id": "runtime_delivery_package",
                        "capability": "validate_runtime_delivery",
                        "depends_on": ["model_contract"],
                        "approval_checkpoint_required": True,
                        "inputs": {
                            "package_dir": package.package_dir.as_posix(),
                            "output_path": (
                                graph_dir / "runtime_delivery_validation.json"
                            ).as_posix(),
                        },
                    },
                ],
            },
        )
        graph = run_local_task_graph_fixture(graph_path, graph_dir)

        self.assertTrue(office.complete)
        self.assertTrue(model.complete)
        self.assertTrue(browser.complete)
        self.assertTrue(delivery.complete)
        self.assertTrue(graph.success)
        self.assertEqual(workbook_hash_before, sha256_file(workbook_path))
        self.assertEqual(fixture_hash_before, sha256_file(fixture_path))
        self.assertFalse(read_json(graph.execution_manifest_path)["network_runtime_allowed"])


if __name__ == "__main__":
    unittest.main()
