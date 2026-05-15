import hashlib
import json
import tempfile
import unittest
import uuid
from pathlib import Path

from openpyxl import Workbook
from openpyxl import load_workbook

from kernel.personal_ai.adapters.xlsx_readonly_runtime import inspect_xlsx_readonly
from kernel.personal_ai.adapters.xlsx_output_writer import (
    approve_xlsx_output,
    create_approved_xlsx_output,
    plan_xlsx_output,
    validate_xlsx_output,
)
from kernel.personal_ai.hash_utils import sha256_file


def read_json(path):
    return json.loads(Path(path).read_text(encoding="utf-8"))


def workbook_text_payload(workbook_path):
    loaded = load_workbook(workbook_path, read_only=True, data_only=False)
    try:
        values = []
        properties = loaded.properties
        for value in (
            properties.creator,
            properties.lastModifiedBy,
            properties.title,
            properties.subject,
            properties.description,
        ):
            if value is not None:
                values.append(str(value))
        for sheet in loaded.worksheets:
            values.append(sheet.title)
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value is not None:
                        values.append(str(value))
        return "\n".join(values)
    finally:
        loaded.close()


class XlsxOutputWriterTests(unittest.TestCase):
    def build_planned_workspace(self):
        temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(temp_dir.cleanup)
        root = Path(temp_dir.name)
        input_dir = root / "input"
        inspection_dir = root / "inspection"
        output_dir = root / "output"
        input_dir.mkdir()
        inspection_dir.mkdir()
        output_dir.mkdir()
        workbook_path = input_dir / "source.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Data"
        sheet["A1"] = "RAW_HEADER_SECRET"
        sheet["A2"] = "RAW_CELL_SECRET"
        workbook.save(workbook_path)
        inspection = inspect_xlsx_readonly(workbook_path, inspection_dir)
        plan = plan_xlsx_output(
            workbook_path,
            inspection.xlsx_inspection_path,
            output_dir,
        )
        return root, input_dir, workbook_path, inspection, output_dir, plan

    def build_workspace(self):
        root, input_dir, workbook_path, inspection, output_dir, plan = (
            self.build_planned_workspace()
        )
        approval_path = root / "approval.json"
        approval = approve_xlsx_output(
            plan.plan_path,
            approval_path,
            approved=True,
            human_reviewed=True,
            reviewer_id="reviewer-001",
        )
        return root, input_dir, workbook_path, inspection, output_dir, plan, approval

    def test_plan_requires_hash_bound_inspection(self):
        _, _, workbook_path, inspection, output_dir, _, _ = self.build_workspace()
        second_output_dir = output_dir.parent / "second"
        second_output_dir.mkdir()

        result = plan_xlsx_output(
            workbook_path,
            inspection.xlsx_inspection_path,
            second_output_dir,
        )
        plan = read_json(result.plan_path)

        self.assertEqual(plan["input_sha256"], sha256_file(workbook_path))
        self.assertEqual(
            plan["xlsx_inspection_sha256"],
            sha256_file(inspection.xlsx_inspection_path),
        )
        self.assertTrue(plan["required_human_approval"])

    def test_plan_can_redact_input_workbook_path_in_audit_artifact(self):
        _, input_dir, workbook_path, inspection, output_dir, _, _ = self.build_workspace()
        second_output_dir = output_dir.parent / "redacted-plan"
        second_output_dir.mkdir()

        result = plan_xlsx_output(
            workbook_path,
            inspection.xlsx_inspection_path,
            second_output_dir,
            redact_input_path=True,
        )
        plan = read_json(result.plan_path)
        plan_text = result.plan_path.read_text(encoding="utf-8")

        self.assertEqual(plan["input_workbook_path"], "[redacted-input-path]")
        self.assertTrue(plan["input_workbook_path_redacted"])
        self.assertNotIn(input_dir.as_posix(), plan_text)

    def test_rejects_malformed_inspection_json_schema(self):
        _, _, workbook_path, _, output_dir, _, _ = self.build_workspace()
        second_output_dir = output_dir.parent / "malformed-inspection"
        second_output_dir.mkdir()
        bad_inspection_path = output_dir.parent / "bad-inspection.json"
        bad_inspection = {
            "inspection_type": "personal_ai_execution_os_v2_xlsx_readonly_inspection",
            "input_workbook": {"sha256": sha256_file(workbook_path)},
            "sheet_count": 1,
            "raw_cell_values_copied": False,
            "required_human_approval": True,
        }
        bad_inspection_path.write_text(
            json.dumps(bad_inspection, sort_keys=True), encoding="utf-8"
        )

        with self.assertRaisesRegex(ValueError, "sheets is malformed"):
            plan_xlsx_output(workbook_path, bad_inspection_path, second_output_dir)

    def test_approval_required(self):
        _, _, workbook_path, _, output_dir, plan, _ = self.build_workspace()
        missing_approval_path = output_dir.parent / "missing_approval.json"

        with self.assertRaises(ValueError):
            create_approved_xlsx_output(
                workbook_path,
                plan.plan_path,
                missing_approval_path,
                output_dir,
            )

    def test_approve_xlsx_output_requires_explicit_approved(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(ValueError, "explicit approved true is required"):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                human_reviewed=True,
                reviewer_id="reviewer-001",
            )

    def test_approve_xlsx_output_requires_explicit_human_reviewed(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(
            ValueError, "explicit human_reviewed true is required"
        ):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                approved=True,
                reviewer_id="reviewer-001",
            )

    def test_approve_xlsx_output_requires_explicit_reviewer_id(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(ValueError, "explicit reviewer_id is required"):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                approved=True,
                human_reviewed=True,
            )

    def test_approve_xlsx_output_rejects_blank_reviewer_id(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(ValueError, "explicit reviewer_id is required"):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                approved=True,
                human_reviewed=True,
                reviewer_id="  ",
            )

    def test_approve_xlsx_output_rejects_placeholder_reviewer_id(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        for reviewer_id in ("local_human_review", "default", "anonymous"):
            with self.subTest(reviewer_id=reviewer_id):
                with self.assertRaisesRegex(
                    ValueError, "placeholder reviewer_id is not allowed"
                ):
                    approve_xlsx_output(
                        plan.plan_path,
                        root / f"approval-{reviewer_id}.json",
                        approved=True,
                        human_reviewed=True,
                        reviewer_id=reviewer_id,
                    )

    def test_approve_xlsx_output_rejects_approved_false(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(ValueError, "explicit approved true is required"):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                approved=False,
                human_reviewed=True,
                reviewer_id="reviewer-001",
            )

    def test_approve_xlsx_output_rejects_human_reviewed_false(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        with self.assertRaisesRegex(
            ValueError, "explicit human_reviewed true is required"
        ):
            approve_xlsx_output(
                plan.plan_path,
                root / "approval.json",
                approved=True,
                human_reviewed=False,
                reviewer_id="reviewer-001",
            )

    def test_approve_xlsx_output_accepts_explicit_reviewer_approval(self):
        root, _, _, _, _, plan = self.build_planned_workspace()

        result = approve_xlsx_output(
            plan.plan_path,
            root / "approval.json",
            approved=True,
            human_reviewed=True,
            reviewer_id=" reviewer-001 ",
        )
        approval = read_json(result.approval_path)

        self.assertTrue(result.approved)
        self.assertTrue(approval["approved"])
        self.assertTrue(approval["human_reviewed"])
        self.assertEqual(approval["reviewer_id"], "reviewer-001")

    def test_hash_mismatch_rejected(self):
        _, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()
        workbook = Workbook()
        workbook.active["A1"] = "changed"
        workbook.save(workbook_path)

        with self.assertRaises(ValueError):
            create_approved_xlsx_output(
                workbook_path,
                plan.plan_path,
                approval.approval_path,
                output_dir,
            )

    def test_creates_output_manifest_hash_summary_and_validation(self):
        _, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()

        result = create_approved_xlsx_output(
            workbook_path,
            plan.plan_path,
            approval.approval_path,
            output_dir,
        )
        manifest = read_json(result.output_manifest_path)
        validation = read_json(result.validation_report_path)

        self.assertTrue(result.output_workbook_path.exists())
        self.assertTrue(result.output_manifest_path.exists())
        self.assertTrue(result.delivery_summary_path.exists())
        self.assertTrue(result.validation_report_path.exists())
        self.assertEqual(manifest["output_workbook_sha256"], result.output_sha256)
        self.assertEqual(
            manifest["output_workbook_metadata_policy"]["creator"],
            "personal_ai_local_runtime",
        )
        self.assertTrue(validation["complete"])
        self.assertTrue(validation["manifest_hash_verified"])

        loaded = load_workbook(result.output_workbook_path, read_only=True)
        try:
            self.assertEqual(loaded.properties.creator, "personal_ai_local_runtime")
            self.assertEqual(
                loaded.properties.lastModifiedBy, "personal_ai_local_runtime"
            )
        finally:
            loaded.close()

    def test_refuses_existing_output_workbook(self):
        _, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()
        (output_dir / "derived_xlsx_summary.xlsx").write_text("", encoding="utf-8")

        with self.assertRaises(ValueError):
            create_approved_xlsx_output(
                workbook_path,
                plan.plan_path,
                approval.approval_path,
                output_dir,
            )

    def test_rejects_path_traversal_output_workbook_name(self):
        _, _, workbook_path, inspection, output_dir, _, _ = self.build_workspace()
        second_output_dir = output_dir.parent / "path-traversal"
        second_output_dir.mkdir()

        with self.assertRaisesRegex(ValueError, "must be a filename"):
            plan_xlsx_output(
                workbook_path,
                inspection.xlsx_inspection_path,
                second_output_dir,
                output_workbook_name="../escape.xlsx",
            )

    def test_preserves_input_and_does_not_leak_raw_values_to_audit_artifacts(self):
        _, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()
        before_hash = sha256_file(workbook_path)

        result = create_approved_xlsx_output(
            workbook_path,
            plan.plan_path,
            approval.approval_path,
            output_dir,
        )

        self.assertEqual(sha256_file(workbook_path), before_hash)
        for path in (
            plan.plan_path,
            approval.approval_path,
            result.output_manifest_path,
            result.delivery_summary_path,
            result.validation_report_path,
        ):
            text = path.read_text(encoding="utf-8")
            self.assertNotIn("RAW_HEADER_SECRET", text)
            self.assertNotIn("RAW_CELL_SECRET", text)

    def test_dynamic_sentinel_values_do_not_leak_to_json_or_markdown_artifacts(self):
        root, input_dir, workbook_path, _, output_dir, _, _ = self.build_workspace()
        secret_sheet = "S" + uuid.uuid4().hex[:20]
        secret_sheet_sha256 = hashlib.sha256(
            secret_sheet.encode("utf-8")
        ).hexdigest()
        secret_header = "header_" + uuid.uuid4().hex
        secret_cell = "cell_" + uuid.uuid4().hex
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = secret_sheet
        sheet["A1"] = secret_header
        sheet["A2"] = secret_cell
        workbook.save(workbook_path)
        inspection_dir = root / "redacted-inspection"
        plan_dir = root / "redacted-output"
        inspection_dir.mkdir()
        plan_dir.mkdir()
        inspection = inspect_xlsx_readonly(
            workbook_path,
            inspection_dir,
            redact_sheet_names=True,
            redact_input_path=True,
        )
        plan = plan_xlsx_output(
            workbook_path,
            inspection.xlsx_inspection_path,
            plan_dir,
            redact_input_path=True,
        )
        approval = approve_xlsx_output(
            plan.plan_path,
            root / "redacted-approval.json",
            approved=True,
            human_reviewed=True,
            reviewer_id="reviewer-002",
        )

        result = create_approved_xlsx_output(
            workbook_path,
            plan.plan_path,
            approval.approval_path,
            plan_dir,
        )

        forbidden_values = (
            secret_sheet,
            secret_sheet_sha256,
            secret_header,
            secret_cell,
            workbook_path.as_posix(),
            input_dir.as_posix(),
        )
        workbook_payload = workbook_text_payload(result.output_workbook_path)
        for sentinel in forbidden_values:
            self.assertNotIn(sentinel, workbook_payload)

        audit_paths = (
            sorted(inspection_dir.iterdir())
            + sorted(plan_dir.iterdir())
            + [approval.approval_path]
        )
        for path in audit_paths:
            if path.suffix.lower() not in (".json", ".md"):
                continue
            text = path.read_text(encoding="utf-8")
            for sentinel in forbidden_values:
                self.assertNotIn(sentinel, text)
        self.assertTrue(result.complete)

    def test_validate_xlsx_output_writes_separate_report_without_overwrite(self):
        root, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()
        create_approved_xlsx_output(
            workbook_path,
            plan.plan_path,
            approval.approval_path,
            output_dir,
        )
        report_path = root / "validation-rerun.json"

        validation = validate_xlsx_output(output_dir, report_path)

        self.assertTrue(validation.complete)
        self.assertTrue(report_path.exists())
        with self.assertRaises(ValueError):
            validate_xlsx_output(output_dir, report_path)

    def test_validate_xlsx_output_rejects_malformed_manifest_schema(self):
        root, _, workbook_path, _, output_dir, plan, approval = self.build_workspace()
        result = create_approved_xlsx_output(
            workbook_path,
            plan.plan_path,
            approval.approval_path,
            output_dir,
        )
        manifest = read_json(result.output_manifest_path)
        manifest.pop("output_workbook_metadata_policy")
        result.output_manifest_path.write_text(
            json.dumps(manifest, sort_keys=True), encoding="utf-8"
        )

        with self.assertRaisesRegex(ValueError, "metadata policy is malformed"):
            validate_xlsx_output(output_dir, root / "bad-manifest-validation.json")


if __name__ == "__main__":
    unittest.main()
